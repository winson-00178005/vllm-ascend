# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def add_title_slide(prs, title, subtitle):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    sub_shape = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(1))
    tf = sub_shape.text_frame
    p = tf.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(24)
    p.alignment = PP_ALIGN.CENTER
    return slide

def add_content_slide(prs, title, bullets):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    content_shape = slide.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(8.6), Inches(5))
    tf = content_shape.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = "• " + bullet
        p.font.size = Pt(20)
        p.space_after = Pt(12)
    return slide

def add_two_column_slide(prs, title, left_title, left_bullets, right_title, right_bullets):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True

    left_title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(4.3), Inches(0.5))
    tf = left_title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = left_title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 102, 204)

    left_content_shape = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(4.3), Inches(4.5))
    tf = left_content_shape.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(left_bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = "• " + bullet
        p.font.size = Pt(16)
        p.space_after = Pt(8)

    right_title_shape = slide.shapes.add_textbox(Inches(5.2), Inches(1.2), Inches(4.3), Inches(0.5))
    tf = right_title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = right_title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 153, 76)

    right_content_shape = slide.shapes.add_textbox(Inches(5.2), Inches(1.8), Inches(4.3), Inches(4.5))
    tf = right_content_shape.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(right_bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = "• " + bullet
        p.font.size = Pt(16)
        p.space_after = Pt(8)
    return slide

def add_table_slide(prs, title, headers, rows):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True

    cols = len(headers)
    table = slide.shapes.add_table(len(rows) + 1, cols, Inches(0.5), Inches(1.3), Inches(9), Inches(4)).table
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(14)
    for row_idx, row in enumerate(rows):
        for col_idx, cell_text in enumerate(row):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = cell_text
            cell.text_frame.paragraphs[0].font.size = Pt(12)
    return slide

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs,
        "vLLM Ascend ST 测试框架设计方案",
        "基于场景化的系统测试框架 | 2026年4月")

    add_content_slide(prs, "目录", [
        "一、现有测试方案的痛点",
        "二、新框架的设计目标",
        "三、框架核心能力",
        "四、目录结构与迁移策略",
        "五、收益与价值",
        "六、下一步计划"
    ])

    add_content_slide(prs, "一、现有测试方案的痛点", [
        "【场景切换困难】单卡/多卡/310P 测试代码分散，环境变量硬编码",
        "【模型配置散落】模型列表分散在各个测试文件中，修改困难",
        "【条件执行繁琐】使用 @pytest.mark.skipif 手工维护跳过条件",
        "【测试复用性差】相似测试逻辑在多个文件中重复",
        "【迁移成本高】现有测试无法快速适配新场景，需大量改写"
    ])

    add_content_slide(prs, "二、新框架的设计目标", [
        "【场景加载机制】支持 SINGLECARD / MULTICARD_2Cards / MULTICARD_4Cards / 310P_SINGLECARD",
        "【模型配置化】通过 YAML 文件统一管理支持的模型列表",
        "【条件执行装饰器】@require_scene / @require_model 声明所需场景/模型",
        "【平滑迁移】现有 @pytest.mark.parametrize 模式零改动迁移",
        "【环境隔离】自动保存/恢复环境变量，测试间互不干扰"
    ])

    add_two_column_slide(prs, "三、框架核心能力",
        "场景管理 (SceneManager)",
        [
            "单例模式，线程安全",
            "支持场景动态切换",
            "自动保存/恢复环境变量",
            "场景配置示例：",
            "  - SINGLECARD: tp_size=1",
            "  - MULTICARD_2Cards: tp_size=2",
            "  - MULTICARD_4Cards: tp_size=4"
        ],
        "装饰器 (Decorators)",
        [
            "@require_scene - 场景级别筛选",
            "@require_model - 模型级别筛选",
            "@require_scene_and_model - 组合筛选",
            "scene_aware_parametrize - 兼容层",
            "自动跳过不匹配的测试"
        ])

    add_content_slide(prs, "四、目录结构", [
        "tests/st/                          # 新建 ST 框架（与 e2e/ut 平级）",
        "├── config/",
        "│   ├── scene.yaml                 # 场景配置",
        "│   └── models.yaml                # 模型列表",
        "├── framework/",
        "│   ├── scene_manager.py           # 场景管理器",
        "│   ├── model_config.py            # 模型配置加载器",
        "│   ├── decorators.py             # 条件执行装饰器",
        "│   ├── compat.py                  # 兼容层（关键！）",
        "│   └── extensions.py              # 性能监控、精度对比",
        "└── testcases/                    # 测试用例"
    ])

    add_content_slide(prs, "四、迁移策略：渐进式切换", [
        "【零侵入】e2e/ 和 ut/ 目录保持不变",
        "【逐个迁移】测试用例可逐个迁移，不需要一次性全部改完",
        "【装饰器可选】不迁移的测试也能正常运行",
        "【VllmRunner 兼容】现有 with VllmRunner(...) 模式保持不变",
        "【复用 Fixtures】新框架 fixtures 不影响现有 vl_config 等"
    ])

    add_table_slide(prs, "五、收益与价值", [
        "收益项", "现状", "新框架"
    ], [
        ["测试执行效率", "手工切换环境变量", "自动根据场景执行"],
        ["模型配置管理", "分散在各测试文件", "统一 YAML 管理"],
        ["条件执行维护", "大量 @skipif 散落", "装饰器声明式"],
        ["代码复用性", "重复代码多", "框架提供统一抽象"],
        ["测试隔离性", "环境变量易污染", "自动保存/恢复"],
        ["迁移成本", "需大量改写", "渐进式、零侵入"]
    ])

    add_content_slide(prs, "五、框架扩展能力（基于 PR #8557）", [
        "【PerformanceMonitor】性能基准测试，支持 TPS / Latency 指标收集",
        "【PrecisionComparator】精度对比，自动对比新旧实现输出差异",
        "【Version Checker】vllm_version_is() 版本兼容性检查",
        "【VllmRunner 封装】简化推理调用，降低测试用例编写门槛"
    ])

    add_content_slide(prs, "六、下一步计划", [
        "【第一阶段】框架验证（当前）",
        "  - 完成框架自测试，验证基本功能可用",
        "  - 在 NPU 环境进行实际测试",
        "【第二阶段】试点迁移",
        "  - 选择 2-3 个典型测试进行迁移",
        "  - 验证迁移成本低、可回滚",
        "【第三阶段】全面推广",
        "  - 逐步将 e2e 测试迁移到新框架",
        "  - 建立场景-模型矩阵报告"
    ])

    add_content_slide(prs, "总结", [
        "ST 测试框架是 vLLM Ascend 测试体系的重要补充",
        "通过场景化设计，解决多卡/多环境测试的复杂性",
        "配置化 + 装饰器模式，降低测试编写门槛",
        "渐进式迁移策略，保证现有测试稳定性的同时完成升级",
        "当前状态：框架代码已完成，验证中",
        "",
        "代码位置：tests/st/ | 文档：docs/st_framework_design.md"
    ])

    prs.save("docs/st_framework_presentation.pptx")
    print("PPT 生成成功: docs/st_framework_presentation.pptx")

if __name__ == "__main__":
    create_presentation()
